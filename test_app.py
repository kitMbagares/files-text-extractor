"""
test_app.py — Unit and integration tests for Extract-Kit.

Test structure:
  - Unit tests:        Test individual functions in isolation
  - Integration tests: Test the full /upload/ endpoint via HTTP

All test files (PDF, DOCX, XLSX, CSV, images) are generated in memory
so no external fixture files are needed.

Run: python -m pytest test_app.py -v
"""

import io
import csv
import pytest
from unittest.mock import patch
from httpx import AsyncClient, ASGITransport
from fastapi import HTTPException
from PIL import Image
import fitz
from docx import Document
from openpyxl import Workbook

# Patch LanguageTool before importing app (it requires Java + 200MB download)
with patch("language_tool_python.LanguageTool"):
    from app import app
    from validators import detect_file_type, validate_file_size, validate_mode_for_type
    from extractors.pdf import (
        extract_text_from_pdf,
        extract_tables_from_pdf,
        extract_images_from_pdf,
        get_pdf_page_count,
    )
    from extractors.docx_ext import extract_text_from_docx
    from extractors.spreadsheet import extract_data_from_xlsx, extract_data_from_csv
    from text_utils import flatten_table_to_text

API_KEY = "test-key"


# ═══════════════════════════════════════════════════════════════════
#  Test file generators
# ═══════════════════════════════════════════════════════════════════
# These create minimal valid files in memory for each format.


def make_pdf(text="Hello World"):
    """Create a single-page PDF with the given text."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_pdf_with_image():
    """Create a PDF with a small embedded red PNG image."""
    doc = fitz.open()
    page = doc.new_page()
    img = Image.new("RGB", (10, 10), color="red")
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    page.insert_image(fitz.Rect(72, 72, 172, 172), stream=img_buf.getvalue())
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_image():
    """Create a simple white PNG image."""
    img = Image.new("RGB", (200, 50), color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def make_docx(text="Hello from Word"):
    """Create a .docx with a single paragraph."""
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx(data=None):
    """Create a .xlsx with one sheet of data."""
    if data is None:
        data = [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in data:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv(rows=None):
    """Create a CSV as UTF-8 bytes."""
    if rows is None:
        rows = [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


# ═══════════════════════════════════════════════════════════════════
#  Unit tests: validators.py
# ═══════════════════════════════════════════════════════════════════


class TestDetectFileType:
    """Test file extension -> type mapping."""

    def test_pdf(self):
        assert detect_file_type("doc.pdf") == "pdf"

    def test_image_formats(self):
        for ext in ("png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp"):
            assert detect_file_type(f"photo.{ext}") == "image"

    def test_docx(self):
        assert detect_file_type("report.docx") == "docx"

    def test_xlsx(self):
        assert detect_file_type("data.xlsx") == "xlsx"

    def test_csv(self):
        assert detect_file_type("data.csv") == "csv"

    def test_unsupported_extension_returns_400(self):
        with pytest.raises(HTTPException) as exc_info:
            detect_file_type("file.xyz")
        assert exc_info.value.status_code == 400
        assert ".xyz" in exc_info.value.detail

    def test_no_extension_returns_400(self):
        with pytest.raises(HTTPException) as exc_info:
            detect_file_type("noextension")
        assert exc_info.value.status_code == 400

    def test_empty_filename_returns_400(self):
        with pytest.raises(HTTPException) as exc_info:
            detect_file_type("")
        assert exc_info.value.status_code == 400

    def test_case_insensitive(self):
        assert detect_file_type("DOC.PDF") == "pdf"
        assert detect_file_type("photo.JPG") == "image"


class TestValidateFileSize:
    """Test empty file and oversized file rejection."""

    def test_valid_size_passes(self):
        validate_file_size(b"some content", "test.pdf")  # should not raise

    def test_empty_file_returns_400(self):
        with pytest.raises(HTTPException) as exc_info:
            validate_file_size(b"", "test.pdf")
        assert exc_info.value.status_code == 400
        assert "empty" in exc_info.value.detail.lower()

    def test_oversized_file_returns_413(self):
        huge = b"x" * (50 * 1024 * 1024 + 1)
        with pytest.raises(HTTPException) as exc_info:
            validate_file_size(huge, "big.pdf")
        assert exc_info.value.status_code == 413
        assert "too large" in exc_info.value.detail.lower()


class TestValidateModeForType:
    """Test that invalid mode + file type combos are rejected."""

    def test_pdf_accepts_all_modes(self):
        for mode in ("text", "tables", "images", "ocr", "full"):
            validate_mode_for_type(mode, "pdf")  # should not raise

    def test_image_accepts_text_only(self):
        validate_mode_for_type("text", "image")

    def test_image_rejects_tables(self):
        with pytest.raises(HTTPException) as exc_info:
            validate_mode_for_type("tables", "image")
        assert exc_info.value.status_code == 400
        assert "tables" in exc_info.value.detail

    def test_docx_rejects_images_mode(self):
        with pytest.raises(HTTPException) as exc_info:
            validate_mode_for_type("images", "docx")
        assert exc_info.value.status_code == 400

    def test_xlsx_rejects_ocr_mode(self):
        with pytest.raises(HTTPException) as exc_info:
            validate_mode_for_type("ocr", "xlsx")
        assert exc_info.value.status_code == 400


# ═══════════════════════════════════════════════════════════════════
#  Unit tests: extractors/pdf.py
# ═══════════════════════════════════════════════════════════════════


class TestExtractTextFromPdf:
    """Test PDF text extraction via pdfplumber."""

    def test_extracts_text(self):
        result = extract_text_from_pdf(make_pdf("Test content here"))
        assert "Test content here" in result

    def test_empty_pdf_returns_empty(self):
        doc = fitz.open()
        doc.new_page()
        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        result = extract_text_from_pdf(buf.getvalue())
        assert result.strip() == ""

    def test_corrupt_pdf_raises_422(self):
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_pdf(b"not a pdf")
        assert exc_info.value.status_code == 422
        assert "Failed to read PDF" in exc_info.value.detail


class TestExtractTablesFromPdf:
    """Test PDF table extraction."""

    def test_returns_list(self):
        result = extract_tables_from_pdf(make_pdf("Just text"))
        assert isinstance(result, list)

    def test_corrupt_pdf_raises_422(self):
        with pytest.raises(HTTPException) as exc_info:
            extract_tables_from_pdf(b"not a pdf")
        assert exc_info.value.status_code == 422


class TestExtractImagesFromPdf:
    """Test PDF embedded image extraction via PyMuPDF."""

    def test_extracts_embedded_image(self):
        result = extract_images_from_pdf(make_pdf_with_image())
        assert len(result) >= 1
        assert "data_base64" in result[0]
        assert result[0]["page"] == 1

    def test_no_images_returns_empty(self):
        result = extract_images_from_pdf(make_pdf("No images here"))
        assert result == []

    def test_respects_max_images_limit(self):
        result = extract_images_from_pdf(make_pdf_with_image(), max_images=0)
        assert result == []

    def test_corrupt_pdf_raises_422(self):
        with pytest.raises(HTTPException) as exc_info:
            extract_images_from_pdf(b"not a pdf")
        assert exc_info.value.status_code == 422


class TestGetPdfPageCount:
    """Test PDF page counting."""

    def test_single_page(self):
        assert get_pdf_page_count(make_pdf()) == 1

    def test_multi_page(self):
        doc = fitz.open()
        for _ in range(3):
            doc.new_page()
        buf = io.BytesIO()
        doc.save(buf)
        doc.close()
        assert get_pdf_page_count(buf.getvalue()) == 3

    def test_invalid_returns_zero(self):
        assert get_pdf_page_count(b"garbage") == 0


# ═══════════════════════════════════════════════════════════════════
#  Unit tests: extractors/docx_ext.py
# ═══════════════════════════════════════════════════════════════════


class TestExtractTextFromDocx:
    """Test Word document text extraction."""

    def test_extracts_text(self):
        result = extract_text_from_docx(make_docx("Document paragraph text"))
        assert "Document paragraph text" in result

    def test_multiple_paragraphs(self):
        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        buf = io.BytesIO()
        doc.save(buf)
        result = extract_text_from_docx(buf.getvalue())
        assert "First paragraph" in result
        assert "Second paragraph" in result

    def test_corrupt_docx_raises_422(self):
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_docx(b"not a docx")
        assert exc_info.value.status_code == 422
        assert "DOCX" in exc_info.value.detail


# ═══════════════════════════════════════════════════════════════════
#  Unit tests: extractors/spreadsheet.py
# ═══════════════════════════════════════════════════════════════════


class TestExtractDataFromXlsx:
    """Test Excel .xlsx extraction."""

    def test_extracts_rows(self):
        result = extract_data_from_xlsx(make_xlsx([["Name", "Age"], ["Alice", "30"]]))
        assert "Sheet1" in result
        assert len(result["Sheet1"]) == 2
        assert result["Sheet1"][0] == ["Name", "Age"]

    def test_none_cells_become_empty_strings(self):
        result = extract_data_from_xlsx(make_xlsx([["A", None], [None, "B"]]))
        rows = result["Sheet1"]
        assert rows[0] == ["A", ""]
        assert rows[1] == ["", "B"]

    def test_corrupt_xlsx_raises_422(self):
        with pytest.raises(HTTPException) as exc_info:
            extract_data_from_xlsx(b"not xlsx")
        assert exc_info.value.status_code == 422
        assert "XLSX" in exc_info.value.detail


class TestExtractDataFromCsv:
    """Test CSV extraction."""

    def test_extracts_rows(self):
        result = extract_data_from_csv(make_csv([["Name", "Age"], ["Alice", "30"]]))
        assert len(result) == 2
        assert result[0] == ["Name", "Age"]

    def test_latin1_fallback(self):
        text = "Nom,Âge\nAlice,30\n"
        result = extract_data_from_csv(text.encode("latin-1"))
        assert len(result) == 2

    def test_empty_csv(self):
        result = extract_data_from_csv(b"")
        assert result == []


# ═══════════════════════════════════════════════════════════════════
#  Unit tests: text_utils.py
# ═══════════════════════════════════════════════════════════════════


class TestFlattenTableToText:
    """Test table data -> readable text conversion."""

    def test_dict_data_xlsx_format(self):
        data = {"Sheet1": [["A", "B"], ["1", "2"]]}
        result = flatten_table_to_text(data)
        assert "Sheet: Sheet1" in result
        assert "A | B" in result

    def test_list_of_dicts_pdf_table_format(self):
        data = [{"page": 1, "table_index": 0, "rows": [["X", "Y"]]}]
        result = flatten_table_to_text(data)
        assert "Table (page 1)" in result
        assert "X | Y" in result

    def test_list_of_rows_csv_format(self):
        data = [["Col1", "Col2"], ["Val1", "Val2"]]
        result = flatten_table_to_text(data)
        assert "Col1 | Col2" in result

    def test_empty_data(self):
        assert flatten_table_to_text({}) == ""
        assert flatten_table_to_text([]) == ""


# ═══════════════════════════════════════════════════════════════════
#  Integration tests: API endpoints
# ═══════════════════════════════════════════════════════════════════


@pytest.fixture
def auth_header():
    """Valid auth header for test requests."""
    return {"Authorization": f"Bearer {API_KEY}"}


@pytest.mark.anyio
class TestHealthEndpoint:
    """Test GET / health check."""

    async def test_returns_status_and_capabilities(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.get("/")
        assert resp.status_code == 200
        data = resp.json()
        assert data["status"] == "ok"
        assert "version" in data
        assert "ocr_available" in data
        assert "supported_types" in data


@pytest.mark.anyio
class TestUploadEndpoint:
    """Test POST /upload/ with various file types, modes, and error cases."""

    @pytest.fixture(autouse=True)
    def _patch_api_key(self):
        """Use a test API key for all upload tests."""
        with patch("validators.API_KEY", API_KEY):
            yield

    async def _upload(self, filename, content, headers, params=None):
        """Helper to POST a file to /upload/."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            files = {"file": (filename, content, "application/octet-stream")}
            return await client.post("/upload/", files=files, headers=headers, params=params)

    # ── Auth tests ──────────────────────────────────────────────────

    async def test_missing_auth_returns_401(self):
        resp = await self._upload("test.pdf", make_pdf(), {})
        assert resp.status_code == 401

    async def test_wrong_token_returns_401(self):
        resp = await self._upload("test.pdf", make_pdf(), {"Authorization": "Bearer wrong"})
        assert resp.status_code == 401
        assert "Invalid" in resp.json()["detail"]

    async def test_no_bearer_prefix_returns_401(self):
        resp = await self._upload("test.pdf", make_pdf(), {"Authorization": API_KEY})
        assert resp.status_code == 401

    # ── File validation tests ───────────────────────────────────────

    async def test_empty_file_returns_400(self, auth_header):
        resp = await self._upload("test.pdf", b"", auth_header)
        assert resp.status_code == 400
        assert "empty" in resp.json()["detail"].lower()

    async def test_unsupported_type_returns_400(self, auth_header):
        resp = await self._upload("data.zip", b"fake zip", auth_header)
        assert resp.status_code == 400
        assert ".zip" in resp.json()["detail"]

    async def test_invalid_mode_returns_422(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf(), auth_header, params={"mode": "invalid"})
        assert resp.status_code == 422

    async def test_incompatible_mode_returns_400(self, auth_header):
        resp = await self._upload("test.docx", make_docx(), auth_header, params={"mode": "images"})
        assert resp.status_code == 400
        assert "images" in resp.json()["detail"]

    # ── Corrupt file tests ──────────────────────────────────────────

    async def test_corrupt_pdf_returns_422(self, auth_header):
        resp = await self._upload("bad.pdf", b"not a real pdf", auth_header, params={"correct": "false"})
        assert resp.status_code == 422
        assert "Failed" in resp.json()["detail"]

    async def test_corrupt_docx_returns_422(self, auth_header):
        resp = await self._upload("bad.docx", b"not a real docx", auth_header, params={"correct": "false"})
        assert resp.status_code == 422

    async def test_corrupt_xlsx_returns_422(self, auth_header):
        resp = await self._upload("bad.xlsx", b"not a real xlsx", auth_header, params={"correct": "false"})
        assert resp.status_code == 422

    # ── PDF: text mode ──────────────────────────────────────────────

    async def test_pdf_text_extraction(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf("Integration test"), auth_header, params={"mode": "text", "correct": "false"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["file_type"] == "pdf"
        assert data["mode"] == "text"
        assert "Integration test" in data["extracted_text"]
        assert data["metadata"]["pages"] == 1

    async def test_pdf_default_mode_is_text(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf("Default mode"), auth_header, params={"correct": "false"})
        data = resp.json()
        assert data["mode"] == "text"
        assert "Default mode" in data["extracted_text"]

    # ── PDF: tables mode ────────────────────────────────────────────

    async def test_pdf_table_extraction(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf("Some content"), auth_header, params={"mode": "tables"})
        assert resp.status_code == 200
        assert isinstance(resp.json()["tables"], list)

    # ── PDF: images mode ────────────────────────────────────────────

    async def test_pdf_image_extraction(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf_with_image(), auth_header, params={"mode": "images"})
        assert resp.status_code == 200
        data = resp.json()
        assert len(data["images"]) >= 1
        assert "data_base64" in data["images"][0]

    # ── PDF: full mode ──────────────────────────────────────────────

    async def test_pdf_full_mode_returns_all(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf_with_image(), auth_header, params={"mode": "full", "correct": "false"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["extracted_text"] is not None
        assert data["tables"] is not None
        assert data["images"] is not None

    # ── DOCX ────────────────────────────────────────────────────────

    async def test_docx_text_extraction(self, auth_header):
        resp = await self._upload("test.docx", make_docx("Word doc content"), auth_header, params={"correct": "false"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["file_type"] == "docx"
        assert "Word doc content" in data["extracted_text"]

    # ── XLSX ────────────────────────────────────────────────────────

    async def test_xlsx_data_extraction(self, auth_header):
        resp = await self._upload("test.xlsx", make_xlsx([["Product", "Price"], ["Widget", "9.99"]]), auth_header)
        assert resp.status_code == 200
        data = resp.json()
        assert data["file_type"] == "xlsx"
        assert "Sheet1" in data["tables"]

    # ── CSV ─────────────────────────────────────────────────────────

    async def test_csv_data_extraction(self, auth_header):
        resp = await self._upload("test.csv", make_csv([["City", "Pop"], ["Paris", "2M"]]), auth_header)
        assert resp.status_code == 200
        data = resp.json()
        assert data["file_type"] == "csv"
        assert len(data["tables"]) == 2

    # ── Image (OCR) ─────────────────────────────────────────────────

    async def test_image_ocr_extraction(self, auth_header):
        resp = await self._upload("photo.png", make_image(), auth_header, params={"correct": "false"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["file_type"] == "image"
        assert data["extracted_text"] is not None

    # ── Response envelope ───────────────────────────────────────────

    async def test_response_contains_all_fields(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf("Envelope test"), auth_header, params={"correct": "false"})
        data = resp.json()
        for field in ("filename", "file_type", "mode", "extracted_text", "corrected_text", "tables", "images", "metadata"):
            assert field in data, f"Missing field: {field}"

    async def test_response_includes_processing_time(self, auth_header):
        resp = await self._upload("test.pdf", make_pdf("Timing test"), auth_header, params={"correct": "false"})
        data = resp.json()
        assert "processing_seconds" in data["metadata"]
        assert isinstance(data["metadata"]["processing_seconds"], float)
