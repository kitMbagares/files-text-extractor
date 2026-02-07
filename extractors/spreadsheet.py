"""
extractors/spreadsheet.py — Excel (.xlsx) and CSV extraction.

Reads spreadsheet data into Python lists/dicts for JSON serialization.
  - XLSX: Returns a dict of {sheet_name: [[row], ...]} using openpyxl
  - CSV:  Returns a flat list of [[row], ...] using stdlib csv
"""

import io
import csv
from fastapi import HTTPException
from openpyxl import load_workbook

from config import logger


def extract_data_from_xlsx(file_bytes: bytes) -> dict:
    """
    Extract all sheets from an Excel .xlsx file.

    Returns: {"Sheet1": [["A1", "B1"], ["A2", "B2"]], "Sheet2": [...]}
    None cells are converted to empty strings.
    Raises 422 if the file is not a valid .xlsx.
    """
    try:
        # read_only=True for memory efficiency, data_only=True to get values (not formulas)
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheets = {}

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                # Convert all cells to strings, None becomes ""
                rows.append([str(cell) if cell is not None else "" for cell in row])
            sheets[sheet_name] = rows

        wb.close()
        return sheets
    except Exception as e:
        logger.exception("Error extracting data from xlsx")
        raise HTTPException(status_code=422, detail=f"Failed to read XLSX file: {e}")


def extract_data_from_csv(file_bytes: bytes) -> list:
    """
    Parse a CSV file into a list of rows.

    Returns: [["col1", "col2"], ["val1", "val2"], ...]
    Tries UTF-8 first, falls back to Latin-1 for older files.
    Raises 422 if the file cannot be parsed.
    """
    try:
        # Try UTF-8 first (most common), fall back to Latin-1
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1")

        reader = csv.reader(io.StringIO(text))
        return [row for row in reader]
    except Exception as e:
        logger.exception("Error extracting data from CSV")
        raise HTTPException(status_code=422, detail=f"Failed to read CSV file: {e}")
